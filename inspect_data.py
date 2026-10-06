import openpyxl, pandas as pd, os, glob

PY = "/Users/panjinlong/.workbuddy/binaries/python/envs/default/bin/python3"
files = [
    "/Users/panjinlong/Downloads/AMCExploration_W81K自定义触点.xlsx",
    "/Users/panjinlong/Downloads/AMCExploration_W81K站内外承接路径.xlsx",
    "/Users/panjinlong/Downloads/AMCExploration_W81K分时情况.xlsx",
    "/Users/panjinlong/Downloads/AMCExploration_W81K曝光频次.xlsx",
    "/Users/panjinlong/Downloads/AMCExploration_W81K漏斗分析.xlsx",
    "/Users/panjinlong/Downloads/Campaign_Sep_1_2026.csv",
]

def inspect_xlsx(path):
    print("="*90)
    print("XLSX:", os.path.basename(path))
    print("="*90)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"\n--- SHEET: {ws.title} | max_row={ws.max_row} max_col={ws.max_column}")
        rows = list(ws.iter_rows(values_only=True))
        # print first 14 rows raw to detect headers/metadata
        for i, r in enumerate(rows[:14]):
            print(f"  r{i}: {r}")
        if len(rows) > 14:
            print(f"  ... total {len(rows)} rows")

def inspect_csv(path):
    print("="*90)
    print("CSV:", os.path.basename(path))
    print("="*90)
    df = pd.read_csv(path)
    print("shape:", df.shape)
    print("columns:", list(df.columns))
    print(df.head(8).to_string())
    print("dtypes:\n", df.dtypes)

for f in files:
    print("\n\n############################################")
    if f.endswith(".csv"):
        inspect_csv(f)
    else:
        inspect_xlsx(f)
