from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = [
    "PingFang SC",
    "Hiragino Sans GB",
    "STHeiti",
    "Heiti SC",
    "Songti SC",
    "Microsoft YaHei",
    "SimSun",
    "Arial Unicode MS",
]

doc = Document()
for font in fonts:
    p = doc.add_paragraph()
    r = p.add_run(f"{font}：中文字体测试｜广告诊断与优化方案")
    r.font.name = font
    r.font.size = Pt(16)
    rf = r._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:ascii"), font)
    rf.set(qn("w:hAnsi"), font)
    rf.set(qn("w:eastAsia"), font)
    rf.set(qn("w:cs"), font)
doc.save("/Users/panjinlong/Documents/agent-master/outputs/W8K4广告诊断-2026-07-28/font_test.docx")
