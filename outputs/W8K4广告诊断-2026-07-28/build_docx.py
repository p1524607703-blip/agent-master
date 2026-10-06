from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "W8K4-SP广告诊断与优化方案.docx"
QA_DIR = ROOT / "docx-qa"
CHART_PATH = QA_DIR / "活动花费与销售额占比.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
GRAY = "5D6670"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "EAF4E3"
PALE_GOLD = "FFF4CC"
RISK_RED = "9B1C1C"
GOLD = "7A5A00"
GREEN = "2E6B3D"
WHITE = "FFFFFF"
BLACK = "111111"

BODY_FONT = "Calibri"
EAST_ASIA_FONT = "PingFang SC"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(
    run,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    name: str = BODY_FONT,
):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="C7CDD4", size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        old_grid.append(col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_cell_text(cell, *, size=8.5, color=BLACK, bold=False, align=None):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        for run in p.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(
    doc,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    *,
    numeric_cols: set[int] | None = None,
    font_size=8.5,
    header_fill=LIGHT_GRAY,
):
    numeric_cols = numeric_cols or set()
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        header.cells[i].text = text
        set_cell_shading(header.cells[i], header_fill)
        format_cell_text(
            header.cells[i],
            size=8.5,
            color=INK,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            if row_idx % 2 == 1:
                set_cell_shading(cells[i], "FAFBFC")
            align = WD_ALIGN_PARAGRAPH.CENTER if i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_text(cells[i], size=font_size, align=align)
    set_table_geometry(table, widths_dxa)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=GRAY)


def add_bottom_border(paragraph, color=BLUE, size=10):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption" in styles:
        caption = styles["Caption"]
        caption.font.name = BODY_FONT
        caption.font.size = Pt(9)
        caption.font.italic = True
        caption.font.color.rgb = rgb(GRAY)
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(8)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = True


def create_numbering(doc: Document, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs, default=0) + 1
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multilevel = OxmlElement("w:multiLevelType")
    multilevel.set(qn("w:val"), "singleLevel")
    abstract.append(multilevel)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def add_list_item(doc, text: str, num_id: int, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_callout(doc, label: str, text: str, *, fill=CALLOUT, label_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=fill, size=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, size=10.5, color=label_color, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_metric_strip(doc):
    values = [
        ("总花费", "$1,076.40"),
        ("广告销售额", "$2,222.09"),
        ("归因购买量", "74"),
        ("整体 ACOS", "48.44%"),
    ]
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [2340, 2340, 2340, 2340], indent_dxa=0)
    set_table_borders(table, color="D7E2EC", size=5)
    for i, (label, value) in enumerate(values):
        cell = table.cell(0, i)
        set_cell_shading(cell, PALE_BLUE if i != 3 else PALE_GOLD)
        set_cell_margins(cell, top=160, bottom=160, start=120, end=120)
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(label)
        set_run_font(r1, size=9, color=GRAY, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(value)
        set_run_font(r2, size=17, color=INK, bold=True)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def load_chart_font(size: int):
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def create_share_chart(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    w, h = 1600, 860
    image = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_chart_font(48)
    label_font = load_chart_font(28)
    small_font = load_chart_font(24)
    value_font = load_chart_font(24)

    draw.text((90, 55), "活动花费与广告销售额占比", fill=f"#{INK}", font=title_font)
    draw.text((90, 120), "单位：账户占比（%）", fill=f"#{GRAY}", font=small_font)

    left, top, right, bottom = 130, 220, 1510, 700
    max_pct = 50
    for pct in range(0, 51, 10):
        y = bottom - (pct / max_pct) * (bottom - top)
        draw.line((left, y, right, y), fill="#D9E0E7", width=2)
        draw.text((55, y - 14), f"{pct}%", fill=f"#{GRAY}", font=small_font)

    campaigns = ["Slip on 手动", "barefoot 手动", "自动低价", "类目低价"]
    spend = [44.96, 37.82, 16.59, 0.63]
    sales = [33.74, 18.89, 46.01, 1.35]
    group_width = (right - left) / len(campaigns)
    bar_width = 90
    colors = ("#2E74B5", "#6A9F58")
    for idx, campaign in enumerate(campaigns):
        center = left + group_width * (idx + 0.5)
        for j, value in enumerate((spend[idx], sales[idx])):
            x0 = center - 105 + j * 120
            x1 = x0 + bar_width
            y0 = bottom - (value / max_pct) * (bottom - top)
            draw.rounded_rectangle((x0, y0, x1, bottom), radius=9, fill=colors[j])
            label = f"{value:.2f}%"
            bbox = draw.textbbox((0, 0), label, font=value_font)
            draw.text(
                (x0 + (bar_width - (bbox[2] - bbox[0])) / 2, y0 - 34),
                label,
                fill=f"#{INK}",
                font=value_font,
            )
        bbox = draw.textbbox((0, 0), campaign, font=label_font)
        draw.text(
            (center - (bbox[2] - bbox[0]) / 2, bottom + 28),
            campaign,
            fill=f"#{BLACK}",
            font=label_font,
        )

    legend_y = 790
    draw.rounded_rectangle((1060, legend_y, 1090, legend_y + 30), radius=4, fill=colors[0])
    draw.text((1105, legend_y - 3), "花费占比", fill=f"#{BLACK}", font=small_font)
    draw.rounded_rectangle((1300, legend_y, 1330, legend_y + 30), radius=4, fill=colors[1])
    draw.text((1345, legend_y - 3), "销售额占比", fill=f"#{BLACK}", font=small_font)
    image.save(path, quality=95)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("AMAZON US · SPONSORED PRODUCTS")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("W8K4 SP 广告诊断与优化方案")
    set_run_font(r, size=25, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("BRONAX 儿童易穿脱帆布宽头赤足鞋｜预算效率、架构与变体归因")
    set_run_font(r, size=13, color=GRAY)

    metadata = [
        ("站点", "Amazon US"),
        ("Listing", "B0D1R1PYGJ"),
        ("分析日期", "2026-07-28"),
        ("数据范围", "W8K4 文件夹内 6 份 SP 原始报表"),
        ("文档状态", "最终复核版"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, size=10.5, color=INK, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=BLACK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(12)
    add_bottom_border(rule, color=BLUE, size=12)


def add_heading(doc, text: str, level: int):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def build_docx():
    QA_DIR.mkdir(parents=True, exist_ok=True)
    create_share_chart(CHART_PATH)

    doc = Document()
    doc.core_properties.title = "W8K4 SP 广告诊断与优化方案"
    doc.core_properties.subject = "Amazon US Sponsored Products 广告数据分析"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Amazon Ads, Sponsored Products, W8K4, BRONAX"
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr1 = hp.add_run("W8K4 SP 广告诊断")
    set_run_font(hr1, size=9, color=GRAY, bold=True)
    hr2 = hp.add_run("    |    Amazon US")
    set_run_font(hr2, size=9, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    bullet_num_id = create_numbering(doc, ordered=False)
    decimal_num_id = create_numbering(doc, ordered=True)

    add_title_block(doc)
    add_heading(doc, "Executive Summary", 1)
    add_callout(
        doc,
        "核心结论",
        "账户累计花费 $1,076.40，带来 74 笔归因购买和 $2,222.09 销售额，整体 ACOS 48.44%。两个高消耗手动广泛活动占 82.78% 花费，仅贡献 52.64% 销售额；自动低价活动以 16.59% 花费贡献 46.01% 销售额。最大的预算黑洞是手动 barefoot 广泛活动及其无法映射到现有定向明细的 $321.79 残差。",
        fill=PALE_BLUE,
    )
    add_metric_strip(doc)
    add_callout(
        doc,
        "决策顺序",
        "先限险，再补数，再收割，最后放量。当前不应给自动活动加预算；应先限制 barefoot 的未知流量风险，并补齐搜索词、广告位、广告商品、完整定向与变更记录报表。",
        fill=PALE_GOLD,
        label_color=GOLD,
    )

    doc.add_page_break()

    add_heading(doc, "1. 数据范围与证据等级", 1)
    p = doc.add_paragraph(
        "本次分析读取 W8K4 文件夹内全部 6 份 CSV，覆盖活动概览、三类手动定向、自动投放组和达成转化商品。所有金额统一为美元，所有 ACOS、ROAS、CVR、CPA 均重新计算并与报表字段交叉核对。"
    )
    p.paragraph_format.space_after = Pt(8)
    for text in [
        "✅ 已验证：六份报表可直接计算或交叉核对的事实。",
        "⚡ 高概率判断：逻辑上最合理，但仍需其他报表验证。",
        "🔵 补数后判断：当前数据不足，不应直接转成投放动作。",
    ]:
        add_list_item(doc, text, bullet_num_id)

    add_callout(
        doc,
        "数据质量提示",
        "活动概览与达成转化商品总计均为 74 笔购买、$2,222.09 销售，但有一笔 $29.99 购买在 barefoot 与自动活动之间归属错位。自动和类目定向明细还存在 $1.28、$0.30 的轻微反向差额，最可能来自导出时点、归因回填或快照异步。",
        fill=CALLOUT,
    )

    add_heading(doc, "2. 账户健康仪表盘", 1)
    doc.add_picture(str(CHART_PATH), width=Inches(6.25))
    cap = doc.add_paragraph(style="Caption")
    cap.add_run("图 1　各活动花费占比与广告销售额占比")

    add_table(
        doc,
        ["活动", "花费", "购买量", "销售额", "ACOS", "ROAS"],
        [
            ["手动 Slip on", "$483.99", "25", "$749.75", "64.55%", "1.55"],
            ["手动 barefoot 广泛", "$407.09", "14", "$419.86", "96.96%", "1.03"],
            ["自动低价", "$178.57", "34", "$1,022.49", "17.46%", "5.73"],
            ["手动类目低价", "$6.75", "1", "$29.99", "22.51%", "4.44"],
            ["合计", "$1,076.40", "74", "$2,222.09", "48.44%", "2.06"],
        ],
        [3100, 1150, 900, 1450, 1250, 1510],
        numeric_cols={1, 2, 3, 4, 5},
        font_size=9,
    )

    add_heading(doc, "2.1 预算配置与产出方向相反", 2)
    p = doc.add_paragraph()
    r1 = p.add_run("事实：")
    set_run_font(r1, bold=True, color=INK)
    r2 = p.add_run(
        "两个手动广泛活动合计花费 $891.08，占 82.78%；仅带来 39 笔购买和 $1,169.61 销售，占销售额 52.64%。自动活动花费占 16.59%，销售额占 46.01%。"
    )
    set_run_font(r2)
    p = doc.add_paragraph()
    r1 = p.add_run("含义：")
    set_run_font(r1, bold=True, color=INK)
    r2 = p.add_run("账户不是缺预算，而是高成本探索层拿走大部分预算，发现层却以更高转化率承接需求。")
    set_run_font(r2)
    p = doc.add_paragraph()
    r1 = p.add_run("动作：")
    set_run_font(r1, bold=True, color=INK)
    r2 = p.add_run("限制低效手动广泛活动继续扩张；自动活动保持预算与竞价，先承担搜索词发现职责。")
    set_run_font(r2)

    add_heading(doc, "3. 预算效率黑洞与数据残差", 1)
    add_table(
        doc,
        ["活动", "活动花费", "已导出定向", "未解释残差", "残差购买", "残差 ACOS"],
        [
            ["手动 Slip on", "$483.99", "$422.18", "$61.81", "2", "103.05%"],
            ["手动 barefoot 广泛", "$407.09", "$85.30", "$321.79", "9", "119.22%"],
            ["自动低价", "$178.57", "$179.85", "-$1.28", "0", "快照差"],
            ["手动类目低价", "$6.75", "$7.05", "-$0.30", "0", "快照差"],
        ],
        [2500, 1200, 1450, 1500, 1150, 1560],
        numeric_cols={1, 2, 3, 4, 5},
        font_size=8.8,
    )
    add_callout(
        doc,
        "最大黑洞",
        "barefoot 已导出的 5 个关键词只解释 $85.30 花费。活动汇总与明细之间仍有 $321.79 花费、260 点击、9 笔购买和 $269.91 销售额，残差 ACOS 为 119.22%。",
        fill="FDECEC",
        label_color=RISK_RED,
    )
    p = doc.add_paragraph(
        "关键词报表已经聚合广泛匹配产生的搜索词，因此不能把残差直接说成“列出的关键词匹配到长尾词”。残差更可能来自未导出的关键词、其他定向类型、状态筛选或快照不一致。正确动作是重新导出全部状态、全部定向类型的 Targeting 报表，在补齐前对活动整体限险。"
    )

    add_heading(doc, "4. 定向层决策", 1)
    add_heading(doc, "4.1 立即限险", 2)
    add_table(
        doc,
        ["对象", "证据", "动作", "观察窗口"],
        [
            [
                "barefoot 活动",
                "ACOS 96.96%；残差 ACOS 119.22%；当前为“提高和降低”",
                "改为“只降低”；日预算 $100 → $25–$30",
                "48 小时看消耗；7 天看购买",
            ],
            [
                "wide kids sneakers",
                "20 点击、$25.75、0 购买",
                "竞价 $0.92 → 约 $0.50；再 10 点击无购买则暂停",
                "下一批 10 点击",
            ],
        ],
        [1600, 2800, 3100, 1860],
        font_size=8.6,
    )

    add_heading(doc, "4.2 保持但不放量", 2)
    for text in [
        "kids natural shoes：13 点击、2 笔购买、ACOS 26.19%。样本仍小，保持当前竞价；等 Search Term 后收割真实成交词。",
        "wide shoes toddler：55 点击、3 笔购买、ACOS 40.24%。保持或小降不超过 10%，先看搜索词。",
        "自动 close-match、loose-match、substitutes：合计 34 笔购买，活动 ACOS 17.46%。保持预算与竞价，避免一次性放大 50%。",
        "类目 3420857011：20 点击、1 笔购买、ACOS 19.77%。至少积累 3 笔购买后再做小步实验。",
    ]:
        add_list_item(doc, text, bullet_num_id)

    add_heading(doc, "4.3 Slip on 两个广泛词", 2)
    add_table(
        doc,
        ["关键词", "点击", "花费", "购买量", "CVR", "ACOS", "当前竞价 / 累计 CPC"],
        [
            ["boys slip on shoes", "207", "$207.97", "13", "6.28%", "53.34%", "$0.50 / $1.00"],
            ["kids slip on shoes", "179", "$214.21", "10", "5.59%", "71.43%", "$0.50 / $1.20"],
        ],
        [2450, 650, 1050, 850, 850, 900, 2610],
        numeric_cols={1, 2, 3, 4, 5, 6},
        font_size=8.5,
    )
    add_callout(
        doc,
        "高概率判断",
        "活动使用“只降低”且当前 Placement 调整为 0 时，当时有效 CPC 不应高于当时竞价。当前 $0.50 竞价与累计 $1.00–$1.20 CPC 的矛盾最可能来自近期降价或快照时点不一致。应冻结竞价 7 天并查 Campaign Change History。",
        fill=CALLOUT,
    )

    add_heading(doc, "5. 转化商品与变体池", 1)
    add_table(
        doc,
        ["指标", "购买量", "占比"],
        [
            ["总归因购买量", "74", "100%"],
            ["光环购买量", "69", "93.24%"],
            ["同父体 B0CPPS6RJ2", "72", "97.30%"],
            ["跨父体", "2", "2.70%"],
            ["被购买商品为 B0D1R1PYGJ", "1", "1.35%"],
        ],
        [5700, 1500, 2160],
        numeric_cols={1, 2},
        font_size=9,
    )
    add_callout(
        doc,
        "正确解释",
        "光环购买量是总购买量的子集，表示用户点击某个广告 ASIN 后购买了另一个 ASIN；不是 74+69=143，也不代表一个点击购买了多双鞋。",
        fill=PALE_GREEN,
        label_color=GREEN,
    )
    p = doc.add_paragraph(
        "97.30% 的归因购买落在同一父体，说明广告对尺码/颜色变体池的销售有明显外溢。但“达成转化的商品”只说明买了什么，不说明广告展示了什么，因此不能据此判断 B0D1R1PYGJ 自身广告转化为零。必须补 Advertised Product 报表，连接广告 ASIN 与被购买 ASIN。"
    )
    p = doc.add_paragraph(
        "管理上应以父体/变体池评估策略与预算，同时用子体库存、Buy Box、尺码、颜色、退货率和贡献利润做控制；SP 实际投放仍需选择合格子 ASIN。"
    )

    add_heading(doc, "6. 产品与 Listing 承接", 1)
    for text in [
        "页面当前售价 $29.99，4.6 星，357 条评价；核心特征是一脚蹬、宽鞋头、零落差、柔软轻量和帆布鞋面。",
        "主定位顺序建议：儿童独立穿脱 → 宽鞋头舒适 → 零落差与灵活鞋底。专业赤足概念作为价值证明，不应盖过家长最直观的购买理由。",
        "尺码偏小是反复出现的购买阻力。应在副图、尺码图、Bullet 与 A+ 中加入脚长测量和选码提示；主图不得加文字。",
        "用柔韧鞋底、宽鞋头、独立穿脱、耐用可清洁和简约帆布外观解释 $29.99，而不是只讲“帆布鞋”。",
        "Search Term 尚未补齐前，不凭活动名称猜测用户搜索意图，也不大幅改写 Listing 关键词。",
    ]:
        add_list_item(doc, text, bullet_num_id)

    add_heading(doc, "7. 分阶段优化计划", 1)
    add_heading(doc, "7.1 0–48 小时：止损与补数", 2)
    for text in [
        "barefoot 从“提高和降低”改为“只降低”，临时压预算至 $25–$30/日。",
        "wide kids sneakers 竞价从 $0.92 降至约 $0.50；再累计 10 点击仍无购买则暂停。",
        "导出完整 Search Term、Placement、Advertised Product、Targeting 和 Campaign Change History。",
        "Slip on 两个广泛词冻结竞价；自动活动不加预算、不提价。",
    ]:
        add_list_item(doc, text, decimal_num_id)

    add_heading(doc, "7.2 3–7 天：重建发现—收割—隔离", 2)
    for text in [
        "从 Search Term 报表提取成交词，按 Slip-on、宽鞋头/赤足、成交 ASIN 三类建立候选精准层。",
        "至少 2–3 笔购买且 CPA 低于目标 CPA 的搜索词，进入 Exact 收割候选；1 笔购买但花费很低的词进入低预算验证组。",
        "新建 Exact 收割后，再把同一搜索词以否定精准加入发现活动；没有搜索词证据不提前加否定。",
        "Discovery 活动使用“只降低”；Harvest 活动可从固定或只降低开始。Placement 加价必须由 Placement 报表证明。",
    ]:
        add_list_item(doc, text, decimal_num_id)

    add_heading(doc, "7.3 8–14 天：统一窗口复评", 2)
    for text in [
        "使用调价后同一 7 天窗口，并等待归因回填后再比较。",
        "有订单目标：新竞价 = 当前竞价 × 目标 ACOS ÷ 实际 ACOS；单次变动限制在 ±15%。",
        "无订单目标：累计无单花费达到目标 CPA 时暂停或降价；目标 CPA 必须来自真实贡献利润。",
        "自动活动只有在实际频繁触顶预算、且高效搜索词仍流失时才增加预算；当前不满足。",
    ]:
        add_list_item(doc, text, decimal_num_id)

    add_heading(doc, "8. 盈亏边界", 1)
    add_callout(
        doc,
        "公式",
        "盈亏平衡 ACOS = 广告前贡献利润率；目标 CPA = 每单广告前贡献利润 × 允许用于获客的比例。",
        fill=PALE_BLUE,
    )
    p = doc.add_paragraph("设置目标 ACOS 与目标 CPA 前，至少补齐：")
    for text in [
        "实际成交价、Coupon 与促销折扣；",
        "Amazon 佣金、FBA 配送与仓储；",
        "产品、包装、头程与落地成本；",
        "退货、退款、销毁或移除损耗；",
        "其他按单变动费用。",
    ]:
        add_list_item(doc, text, bullet_num_id)
    p = doc.add_paragraph(
        "在利润数据缺失时，只能评价相对效率。barefoot 的 96.96% ACOS 意味着它要求约 97% 的广告前贡献利润率才能盈亏平衡，在正常存在平台费用和商品成本的结构下几乎不可能持续。"
    )

    add_heading(doc, "9. 必须补充的数据", 1)
    add_table(
        doc,
        ["缺失数据", "将改变的决策"],
        [
            ["完整 Search Term Report", "识别成交词、浪费词、活动间重叠；决定 Exact 与否定词"],
            ["Placement Report", "区分 Top of Search、Rest of Search、Product Pages 的 CPC/CVR；决定广告位调整"],
            ["Advertised Product Report", "确认哪些子 ASIN 实际消耗预算；连接广告 ASIN 与被购买 ASIN"],
            ["完整 Targeting（全部状态/类型）", "解释 barefoot $321.79 与 Slip on $61.81 残差"],
            ["Campaign Change History + 按日表现", "验证当前竞价与累计 CPC 时点差，建立调价后窗口"],
            ["Budget/Hourly", "确认 barefoot 何时耗尽预算及各时段质量"],
            ["Business Report + Organic Sales", "计算 TACOS、自然单变化与广告增量性"],
            ["子 ASIN 库存/Buy Box/退货率", "防止高效流量落到断货或尺码问题严重的变体"],
            ["子体贡献利润", "设置目标 ACOS、目标 CPA 和无单止损阈值"],
        ],
        [3200, 6160],
        font_size=8.7,
    )

    add_heading(doc, "10. DeepSeek 多轮交叉复核", 1)
    for text in [
        "第一轮独立审计：重算账户与活动指标，识别两个手动广泛活动和 barefoot 残差。",
        "第二轮反向质询：纠正光环占比、竞价/CPC、残差来源和跨活动归因四项错误；DeepSeek 接受 69/74=93.24%、残差 ACOS 119.22% 和一笔 $29.99 归因错位。",
        "第三轮文档教练：采用“事实 → 含义 → 动作 → 观察窗口”的写法，将已验证、高概率与补数后判断分开。",
    ]:
        add_list_item(doc, text, decimal_num_id)
    p = doc.add_paragraph(
        "DeepSeek 第一轮还提出品牌防御、竞品 ASIN、固定竞价和具体 ACOS 阈值，但当前六份报表没有充分证据，本报告未直接采纳。第二轮曾越界推断 B0D1R1PYGJ 自身广告转化为零，最终报告也未采纳。"
    )

    add_heading(doc, "11. 最终判断", 1)
    add_callout(
        doc,
        "最终判断",
        "W8K4 的问题不是单一关键词出价过高，而是预算层、架构层和数据层同时存在硬伤：高消耗手动广泛活动拿走 82.78% 花费；缺少完整的“自动发现 → 精准收割 → 发现层否定隔离”闭环；手动定向明细严重不完整，竞价快照与累计 CPC 还不在同一时点。",
        fill=PALE_BLUE,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("正确顺序：先限险 → 再补数 → 再收割 → 最后放量")
    set_run_font(r, size=14, color=INK, bold=True)

    add_heading(doc, "来源", 2)
    for text in [
        "/Users/panjinlong/Library/Application Support/ziniaobrowserdatas/ziniao browser/欧德思美站/W8K4/",
        "Amazon Listing：https://www.amazon.com/dp/B0D1R1PYGJ?th=1&psc=1",
        "本地复算：outputs/W8K4广告诊断-2026-07-28/analyze_w8k4.py",
        "DeepSeek 复核：OpenCode 会话 ses_05885232dffe8cko5PjD19SxrE",
    ]:
        add_list_item(doc, text, bullet_num_id)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
